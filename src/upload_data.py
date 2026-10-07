"""Parse untrusted uploads in memory. No file writes or shared caches."""
import csv
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
from openpyxl import load_workbook

from src.data_loader import (
    ACTIVE_STATUSES, EDITABLE_STATUSES, ENUMS, REQUIRED_COLUMNS,
    VALUE_INTENTS, DataValidationError,
)

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 100
LABELS = dict(zip(REQUIRED_COLUMNS, (
    '报名编号', '姓名', '身份', '公司 / 项目', '项目阶段', '创业行业 / 关注方向', '报名渠道',
    '是否到场', '对接意向', '当前跟进状态', '跟进负责人', '下一步动作', '备注',
)))
ALIASES = {label: column for column, label in LABELS.items()}
ALIASES.update({
    '编号': 'attendee_id', '参会编号': 'attendee_id', '参会者编号': 'attendee_id',
    '名字': 'name', '参会者姓名': 'name', '参与者姓名': 'name',
    '角色': 'role', '参与者身份': 'role', '参会身份': 'role',
    '公司': 'organization', '机构': 'organization', '公司/项目': 'organization',
    '公司名称': 'organization', '公司/机构': 'organization', '公司 / 机构': 'organization',
    '创业阶段': 'startup_stage', '创业方向': 'field', '关注方向': 'field',
    '方向': 'field', '创业行业': 'field', '行业方向': 'field', '所属行业': 'field',
    '行业': 'field', '渠道': 'registration_channel', '来源渠道': 'registration_channel',
    '报名来源': 'registration_channel', '到场': 'attended', '签到': 'attended',
    '是否签到': 'attended', '到场状态': 'attended', '签到状态': 'attended',
    '对接需求': 'connection_intent', '跟进状态': 'followup_status',
    '负责人': 'followup_owner', '跟进动作': 'next_action',
})
MINIMUM = ('name', 'role', 'registration_channel', 'attended')
ATTENDANCE = {value: '是' for value in ('是', '已到场', '已签到', '到场', '签到', 'yes', 'y', 'true', '1', '1.0')}
ATTENDANCE.update({value: '否' for value in ('否', '未到场', '未签到', '缺席', 'no', 'n', 'false', '0', '0.0')})


def _reject(mask, message):
    if mask.any():
        rows = '、'.join(str(i + 2) for i, bad in enumerate(mask) if bad)[:100]
        raise DataValidationError(f'{message}（表格行 {rows}，含表头）。请修改后重新上传。')


def validate_upload_data(frame):
    """Validate normalized upload records, retaining unknown optional facts as blanks."""
    clean = frame.loc[:, list(REQUIRED_COLUMNS)].fillna('').astype(str)
    for column in clean:
        clean[column] = clean[column].str.strip()
    for column in ('attendee_id', *MINIMUM):
        _reject(clean[column].eq(''), f'「{LABELS[column]}」不能为空')
    _reject(clean.attendee_id.duplicated(keep=False), '报名编号不能重复')
    _reject(~clean.attended.isin(('是', '否')), '是否到场请填写是/否、已签到/未签到、true/false 或 1/0')
    for column, options in [('connection_intent', ENUMS['connection_intent']),
                            ('followup_status', (*EDITABLE_STATUSES, '未到场'))]:
        _reject(~clean[column].isin(('', *options)), f'「{LABELS[column]}」支持留空或：' + '、'.join(options))
    absent = clean.attended.eq('否')
    _reject(absent & (~clean.connection_intent.isin(('', '未到场'))
                     | ~clean.followup_status.isin(('', '未到场'))
                     | ~clean.next_action.isin(('', '未到场')) | clean.followup_owner.ne('')),
            '未到场者的意向、状态、动作只能留空或填写「未到场」，负责人须留空')
    _reject(~absent & clean[['connection_intent', 'followup_status', 'next_action']].eq('未到场').any(axis=1),
            '已到场者不能标记为「未到场」')
    active = clean.followup_status.isin(ACTIVE_STATUSES)
    _reject(active & (clean.followup_owner.eq('') | clean.next_action.isin(('', '暂无', '未到场'))),
            '进行中的跟进须填写负责人和具体下一步动作')
    _reject(clean.followup_status.isin((*ACTIVE_STATUSES, '已完成对接')) & ~clean.connection_intent.isin(VALUE_INTENTS),
            '对接进度须对应明确对接意向')
    terminal = clean.followup_status.isin(('暂不跟进', '已完成对接'))
    _reject(terminal & (clean.followup_owner.ne('') | ~clean.next_action.isin(('', '暂无'))),
            '暂不跟进或已完成对接时，负责人须留空，下一步动作留空或填写「暂无」')
    return clean


def _csv_rows(content):
    for encoding in ('utf-8-sig', 'gb18030'):
        try:
            text = content.decode(encoding)
            break
        except UnicodeError:
            continue
    else:
        raise DataValidationError('无法识别 CSV 编码，请另存为 UTF-8 CSV 后重试。')
    if '\x00' in text:
        raise DataValidationError('CSV 含无法识别的字符，请另存为 UTF-8 CSV。')
    reader = csv.reader(StringIO(text, newline=''), strict=True)
    header = next(reader, [])
    rows = []
    for line, row in enumerate(reader, 2):
        if not row or not any(cell.strip() for cell in row):
            continue
        if len(row) != len(header):
            raise DataValidationError(f'CSV 第 {line} 行列数与表头不一致，请检查逗号与引号。')
        rows.append(row)
        if len(rows) > MAX_ROWS:
            raise DataValidationError(f'最多支持 {MAX_ROWS} 条记录，请拆分文件。')
    return header, rows


def _xlsx_rows(content):
    # Cap expanded archive size before XML parsing; never extract onto disk.
    with ZipFile(BytesIO(content)) as archive:
        if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
            raise DataValidationError('Excel 解压后内容过大，请删除多余工作表或改用 CSV。')
    book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        sheet = book.worksheets[0]
        # Do not trust worksheet dimension metadata to truncate actual data.
        sheet.reset_dimensions()
        rows = []
        for number, cells in enumerate(sheet.iter_rows(), 1):
            if number > MAX_ROWS + 1 or len(cells) > MAX_COLUMNS:
                raise DataValidationError(f'Excel 最多支持 {MAX_ROWS} 行数据、{MAX_COLUMNS} 列，请删除多余行列。')
            if any(cell.data_type in ('f', 'e') for cell in cells):
                raise DataValidationError('Excel 含公式或错误单元格，请复制并粘贴为值后上传。')
            rows.append(['' if cell.value is None else str(cell.value) for cell in cells])
        if not rows:
            return [], []
        width = max(map(len, rows))
        rows = [row + [''] * (width - len(row)) for row in rows]
        # Ignore genuinely empty trailing columns, including formatting-only cells.
        while width and all(not row[width - 1].strip() for row in rows):
            width -= 1
        return rows[0][:width], [row[:width] for row in rows[1:] if any(v.strip() for v in row)]
    finally:
        book.close()


def parse_upload(content: bytes, filename: str) -> pd.DataFrame:
    """Read bytes only, reject ambiguous headers, normalize known Chinese aliases."""
    if not content:
        raise DataValidationError('文件为空，请添加表头和报名记录。')
    if len(content) > MAX_BYTES:
        raise DataValidationError('文件超过 5 MB，请拆分后上传。')
    extension = Path(filename).suffix.lower()
    if extension not in ('.csv', '.xlsx'):
        raise DataValidationError('仅支持 CSV 和 XLSX，请另存为支持的格式。')
    try:
        header, rows = _csv_rows(content) if extension == '.csv' else _xlsx_rows(content)
    except DataValidationError:
        raise
    except Exception as exc:
        # Untrusted parsers raise many backend-specific exceptions. Never expose them in UI.
        raise DataValidationError('文件无法读取，请检查格式是否正确；Excel 请去除密码并另存为 XLSX。') from exc
    if not rows:
        raise DataValidationError('没有报名记录，请在表头下添加数据。')
    if len(header) > MAX_COLUMNS:
        raise DataValidationError(f'最多支持 {MAX_COLUMNS} 列，请删除多余列。')
    columns = [ALIASES.get(str(h).strip(), str(h).strip().lower()) for h in header]
    if '' in columns:
        raise DataValidationError('存在未命名的列，请补充表头或删除空列。')
    if len(set(columns)) != len(columns):
        raise DataValidationError('存在重复字段（可能由中文和英文列名映射到同一字段），请每个字段只保留一列。')
    missing = [LABELS[column] for column in MINIMUM if column not in columns]
    if missing:
        raise DataValidationError('缺少必填字段：' + '、'.join(missing))
    frame = pd.DataFrame(rows, columns=columns).reindex(columns=REQUIRED_COLUMNS, fill_value='')
    frame = frame.fillna('').astype(str).apply(lambda col: col.str.strip())
    existing = set(frame.attendee_id) - {''}
    serial = 1
    for index in frame.index[frame.attendee_id.eq('')]:
        while f'UPLOAD-{serial:06d}' in existing:
            serial += 1
        identifier = f'UPLOAD-{serial:06d}'
        frame.at[index, 'attendee_id'] = identifier
        existing.add(identifier)
    frame['attended'] = frame.attended.map(lambda value: ATTENDANCE.get(value.lower(), value))
    return validate_upload_data(frame)


def template_csv():
    """Header-only template avoids accidentally analysing invented example people."""
    return pd.DataFrame(columns=list(LABELS.values())).to_csv(index=False).encode('utf-8-sig')
