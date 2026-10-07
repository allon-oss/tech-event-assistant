"""Real in-memory CSV/XLSX parsing, validation and upload session contracts."""
from io import BytesIO
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook

from src import data_loader
from src.metrics import calculate_kpis, channel_performance, field_distribution, role_distribution


CSV = '姓名,身份,报名渠道,是否到场\n测试甲,创业者,官网,是\n测试乙,研究员,邮件,否\n'


def xlsx_bytes(rows):
    book = Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class UploadDataTests(unittest.TestCase):
    def parse(self, content=CSV, name='名单.csv'):
        from src.upload_data import parse_upload
        return parse_upload(content.encode('utf-8-sig') if isinstance(content, str) else content, name)

    def test_minimum_chinese_csv_generates_ids_and_leaves_optional_information_blank(self):
        data = self.parse()
        self.assertEqual(data.name.tolist(), ['测试甲', '测试乙'])
        self.assertEqual(data.attendee_id.nunique(), 2)
        self.assertFalse(data.attendee_id.eq('').any())
        for field in ['organization', 'startup_stage', 'field', 'connection_intent',
                      'followup_status', 'followup_owner', 'next_action', 'notes']:
            self.assertEqual(data[field].tolist(), ['', ''])
        self.assertEqual(calculate_kpis(data)['attended'], 1)

    def test_english_csv_and_full_demo_are_supported_without_changes(self):
        english = 'name,role,registration_channel,attended\nTest,其他,网站,true\n'
        self.assertEqual(self.parse(english).attended.tolist(), ['是'])
        source = data_loader.DEFAULT_DATA_PATH.read_bytes()
        pd.testing.assert_frame_equal(self.parse(source), data_loader.load_data())

    def test_xlsx_chinese_and_english_headers_boolean_attendance_and_literal_na(self):
        for headers in [['姓名', '身份', '渠道', '签到'],
                        ['name', 'role', 'registration_channel', 'attended']]:
            with self.subTest(headers=headers):
                data = self.parse(xlsx_bytes([headers, ['NA', '其他', '官网', True],
                                             ['测试乙', '投资人', '邮件', False]]), '名单.XLSX')
                self.assertEqual(data.name.tolist(), ['NA', '测试乙'])
                self.assertEqual(data.attended.tolist(), ['是', '否'])

    def test_attendance_aliases_and_gb18030(self):
        for value, expected in [('1', '是'), ('0', '否'), ('YES', '是'), ('No', '否'),
                                ('已签到', '是'), ('未到场', '否')]:
            with self.subTest(value=value):
                self.assertEqual(self.parse('姓名,身份,报名渠道,是否到场\n甲,其他,网站,' + value).attended[0], expected)
        self.assertEqual(self.parse(CSV.encode('gb18030')).name[0], '测试甲')

    def test_industry_headers_preserve_nontech_fields_in_csv_and_xlsx(self):
        for label in ['field', '方向', '行业', '创业行业', '行业方向', '所属行业', '创业行业 / 关注方向']:
            with self.subTest(label=label):
                headers = ['姓名', '身份', '报名渠道', '是否到场', label]
                row = ['测试店主', '品牌主理人', '商会', '是', '餐饮 / 食品']
                csv_data = self.parse(','.join(headers) + '\n' + ','.join(row))
                excel_data = self.parse(xlsx_bytes([headers, row]), '名单.xlsx')
                self.assertEqual(csv_data.field.tolist(), ['餐饮 / 食品'])
                pd.testing.assert_frame_equal(csv_data, excel_data)
        with self.assertRaises(data_loader.DataValidationError):
            self.parse('姓名,身份,报名渠道,是否到场,方向,创业行业\n甲,创业者,商会,是,餐饮,零售')

    def test_missing_required_columns_and_cells_are_rejected(self):
        for content in ['姓名,身份,报名渠道\n甲,其他,网站',
                        '姓名,身份,报名渠道,是否到场\n ,其他,网站,是',
                        '姓名,身份,报名渠道,是否到场\n甲,,网站,是',
                        '姓名,身份,报名渠道,是否到场\n甲,其他,,是',
                        '姓名,身份,报名渠道,是否到场\n甲,其他,网站,未知']:
            with self.subTest(content=content), self.assertRaises(data_loader.DataValidationError):
                self.parse(content)

    def test_duplicate_mapped_headers_and_duplicate_ids_are_rejected(self):
        for content in ['姓名,name,身份,报名渠道,是否到场\n甲,甲,其他,网站,是',
                        '姓名,姓名,身份,报名渠道,是否到场\n甲,甲,其他,网站,是',
                        'attendee_id,姓名,身份,报名渠道,是否到场\nx,甲,其他,网站,是\nx,乙,其他,网站,否']:
            with self.subTest(content=content), self.assertRaises(data_loader.DataValidationError):
                self.parse(content)
        with self.assertRaises(data_loader.DataValidationError):
            self.parse(xlsx_bytes([['姓名', 'name', '身份', '报名渠道', '是否到场'],
                                   ['甲', '甲', '其他', '网站', '是']]), 'a.xlsx')

    def test_partial_ids_generated_without_collisions(self):
        data = self.parse('attendee_id,姓名,身份,报名渠道,是否到场\nUPLOAD-000001,甲,其他,网站,是\n,乙,其他,网站,否')
        self.assertEqual(data.attendee_id[0], 'UPLOAD-000001')
        self.assertEqual(data.attendee_id.nunique(), 2)

    def test_empty_malformed_unsupported_or_corrupt_files_have_friendly_errors(self):
        for content, name in [(b'', 'a.csv'), (b'bad', 'a.xlsx'), (b'x', 'a.xls'),
                              ('姓名,身份,报名渠道,是否到场', 'a.csv'),
                              (CSV + '甲,其他,网站,是,多一列', 'a.csv')]:
            with self.subTest(name=name, content=content), self.assertRaises(data_loader.DataValidationError):
                self.parse(content, name)

    def test_formula_cells_rejected_and_first_sheet_only(self):
        with self.assertRaises(data_loader.DataValidationError):
            self.parse(xlsx_bytes([['姓名', '身份', '报名渠道', '是否到场'],
                                   ['=1+1', '其他', '网站', '是']]), 'a.xlsx')
        book = Workbook()
        book.active.append(['姓名', '身份', '报名渠道', '是否到场'])
        book.active.append(['首表', '其他', '网站', '是'])
        book.create_sheet('忽略').append(['第二张表不会参与统计'])
        output = BytesIO()
        book.save(output)
        book.close()
        self.assertEqual(self.parse(output.getvalue(), 'a.xlsx').name.tolist(), ['首表'])

    def test_review_distinguishes_missing_intents_and_statuses_from_known_facts(self):
        from src.review_generator import generate_review
        review = generate_review(self.parse())
        self.assertIn('未填写对接意向', review)
        self.assertIn('未填写关注方向', review)
        known_intent = self.parse('姓名,身份,报名渠道,是否到场,对接意向\n甲,其他,官网,是,寻找投资')
        self.assertIn('未填写跟进状态', generate_review(known_intent))

    def test_size_and_row_limits_have_friendly_errors(self):
        from src import upload_data
        with patch.object(upload_data, 'MAX_BYTES', 10), self.assertRaises(data_loader.DataValidationError):
            self.parse()
        with patch.object(upload_data, 'MAX_ROWS', 1), self.assertRaises(data_loader.DataValidationError):
            self.parse()

    def test_unknown_categories_are_counted_and_missing_fields_not_invented(self):
        data = self.parse()
        self.assertEqual(role_distribution(data).set_index('role').loc['研究员', 'count'], 1)
        channels = channel_performance(data).set_index('registration_channel')
        self.assertEqual(channels.loc['官网', 'attendance_rate'], 1)
        self.assertEqual(channels.loc['邮件', 'attended'], 0)
        self.assertEqual(field_distribution(data)['count'].sum(), 0)

    def test_optional_followup_values_preserved_and_contradictions_rejected(self):
        content = '姓名,身份,报名渠道,是否到场,对接意向,当前跟进状态,跟进负责人,下一步动作\n甲,研究员,网站,是,寻找合作伙伴,待跟进,实际负责人,电话沟通'
        data = self.parse(content)
        self.assertEqual(data.followup_owner[0], '实际负责人')
        for invalid in [content.replace('待跟进', '不认识的状态'),
                        content.replace(',是,', ',否,'), content.replace('实际负责人', '')]:
            with self.subTest(invalid=invalid), self.assertRaises(data_loader.DataValidationError):
                self.parse(invalid)

    def test_template_contains_headers_only_and_can_be_filled(self):
        from src.upload_data import template_csv
        template = pd.read_csv(BytesIO(template_csv()), encoding='utf-8-sig')
        self.assertTrue(template.empty)
        self.assertTrue({'姓名', '身份', '报名渠道', '是否到场'}.issubset(template.columns))
        template.loc[0] = [''] * len(template.columns)
        template.loc[0, ['姓名', '身份', '报名渠道', '是否到场']] = ['甲', '其他', '网站', '是']
        self.assertEqual(len(self.parse(template.to_csv(index=False))), 1)


if __name__ == '__main__':
    unittest.main()
